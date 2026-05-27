"""One-off inspection: dump paragraph structure for a sample of .docx files."""
import argparse
import sys
from pathlib import Path

from docx import Document

from ibd_reports.utils.paths import WORKING_REPORTS


def _pick_sample(n: int) -> list[Path]:
    files = sorted(WORKING_REPORTS.glob("*.docx"), key=lambda p: p.stat().st_size)
    if not files:
        print("No .docx files found in WORKING_REPORTS.", file=sys.stderr)
        sys.exit(1)
    half = n // 2
    smallest = files[:half]
    largest = files[-half:]
    seen: set[Path] = set()
    sample: list[Path] = []
    for p in smallest + largest:
        if p not in seen:
            sample.append(p)
            seen.add(p)
    return sample


def _inspect(path: Path) -> None:
    print(f"\n{'='*72}", file=sys.stderr)
    print(f"FILE: {path.name}  ({path.stat().st_size} bytes)", file=sys.stderr)
    print(f"{'='*72}", file=sys.stderr)

    doc = Document(path)

    for i, para in enumerate(doc.paragraphs):
        text = para.text[:200]
        print(f"  [{i:>4}] [{para.style.name}] {text!r}", file=sys.stderr)

    tables = doc.tables
    if not tables:
        print(f"\n  tables: none", file=sys.stderr)
    else:
        first = tables[0]
        rows = len(first.rows)
        cols = len(first.columns)
        print(f"\n  tables: {len(tables)} total; first is {rows}r x {cols}c", file=sys.stderr)


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Dump paragraph structure for a size-stratified sample of .docx files."
    )
    parser.add_argument("--count", type=int, default=6, metavar="N",
                        help="Number of files to inspect (N//2 smallest + N//2 largest).")
    args = parser.parse_args()

    sample = _pick_sample(args.count)
    print(f"Inspecting {len(sample)} files from {WORKING_REPORTS}", file=sys.stderr)

    for path in sample:
        _inspect(path)

    print(f"\n{'='*72}", file=sys.stderr)
    print(f"Done. {len(sample)} files inspected.", file=sys.stderr)


if __name__ == "__main__":
    main()
