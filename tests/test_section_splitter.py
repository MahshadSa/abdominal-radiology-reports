from docx import Document

from ibd_reports.parsing.section_splitter import split_sections


def _make_doc(*texts: str) -> Document:
    doc = Document()
    for p in list(doc.paragraphs):
        p._p.getparent().remove(p._p)
    for text in texts:
        doc.add_paragraph(text)
    return doc


def test_canonical():
    doc = _make_doc(
        "MRI small bowel",
        "Clinical information: Patient has Crohn's disease",
        "Terminal ileum shows wall thickening",
        "Mesentery is normal",
        "Impression: Active disease noted",
        "Sincerely,",
    )
    result = split_sections(doc)
    assert result["modality"] == "MRI small bowel"
    assert result["clinical_info"] == "Patient has Crohn's disease"
    assert result["previous_imaging"] == ""
    assert result["findings"] == "Terminal ileum shows wall thickening\nMesentery is normal"
    assert result["impression"] == "Active disease noted"
    assert result["parse_flags"] == []


def test_comparison_present():
    doc = _make_doc(
        "MRI small bowel",
        "Clinical information: Patient has Crohn's disease",
        "Comparison: 12.03.2021",
        "Terminal ileum shows wall thickening",
        "Impression: Active disease noted",
        "Sincerely,",
    )
    result = split_sections(doc)
    assert result["previous_imaging"] == "12.03.2021"
    assert result["parse_flags"] == []


def test_header_then_content():
    doc = _make_doc(
        "MRI small bowel",
        "Clinical information:",
        "Patient has Crohn's disease",
        "Impression: Active disease noted",
        "Sincerely,",
    )
    result = split_sections(doc)
    assert result["clinical_info"] == "Patient has Crohn's disease"
    assert result["parse_flags"] == []


def test_missing_impression():
    doc = _make_doc(
        "MRI small bowel",
        "Clinical information: Patient has Crohn's disease",
        "Terminal ileum shows wall thickening",
        "Sincerely,",
    )
    result = split_sections(doc)
    assert "no_impression" in result["parse_flags"]
    assert result["impression"] == ""


def test_missing_signoff():
    doc = _make_doc(
        "MRI small bowel",
        "Clinical information: Patient has Crohn's disease",
        "Terminal ileum shows wall thickening",
        "Impression: Active disease noted",
    )
    result = split_sections(doc)
    assert "no_signoff" in result["parse_flags"]
    assert result["findings"] != ""
    assert result["impression"] != ""


def test_case_insensitive():
    doc = _make_doc(
        "MRI small bowel",
        "CLINICAL INFORMATION: Patient has Crohn's disease",
        "Terminal ileum shows wall thickening",
        "Impression: Active disease noted",
        "Sincerely,",
    )
    result = split_sections(doc)
    assert result["clinical_info"] == "Patient has Crohn's disease"
    assert result["parse_flags"] == []


def test_whitespace_tolerance():
    doc = _make_doc(
        "MRI small bowel",
        "Clinical information : Patient has Crohn's disease",
        "Terminal ileum shows wall thickening",
        "Impression: Active disease noted",
        "Sincerely,",
    )
    result = split_sections(doc)
    assert result["clinical_info"] == "Patient has Crohn's disease"
    assert result["parse_flags"] == []


def test_empty_document():
    doc = _make_doc()
    result = split_sections(doc)
    assert result["parse_flags"] == ["empty_document"]
    assert result["modality"] == ""
    assert result["clinical_info"] == ""
    assert result["previous_imaging"] == ""
    assert result["findings"] == ""
    assert result["impression"] == ""


def test_non_ascii_modality():
    doc = _make_doc(
        "MRI small bowel ف",
        "Clinical information: Patient has Crohn's disease",
        "Terminal ileum shows wall thickening",
        "Impression: Active disease noted",
        "Sincerely,",
    )
    result = split_sections(doc)
    assert "non_ascii_in_modality" in result["parse_flags"]
